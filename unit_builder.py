"""语义单元层构建器（unit layer / L3，知识库重构 Phase 1，2026-08-31）。

背景：
- structured（L0/L1/L2）的 L1"段"= 版面列（col），不是内容逻辑单元——名录一列混几十个
  人名碎片，账簿一列里机构/金额/户名交错，导致人读导出全是散字、RAG 检索单元与
  内容单元错位（召回天花板 63%）。
- reassembler.py 的内容类型感知重组（名录条目化、阅读序重排、逐字校验）是最接近
  "内容逻辑"的资产，但只服务学术导出。本模块把该能力提升为管线正中央的 canonical
  衍生层：**每张图一份语义单元文件**，人读导出与 AI 检索语料共享同一套单元。

单元类型（unit_type）：
- roster_entry  名录条目（人名＋科年/籍贯/官职；LLM 聚合零散字词）
- ledger_entry  账目条目（户名＋方向＋币种＋金额原文＋金额归一；规则解析，零 LLM 成本）
- paragraph     连贯段落（prose / letter / document；阅读序重排）
- text_fragment 兜底碎片（账簿页的非账目残余文字等）

数据流：
    data/structured/<stem>.json（canonical，只读）
      ＋ output.xlsx corrected_text（人工校对，经 safe_corrected 粒度校验）
      └→ data/units/<stem>.json（本模块产物，可随时全量重算）

缓存机制（消解 G5 成本问题）：
- 以 source_text_hash（全图行文本 sha1，含校对）为键：未校对过的图一生只算一次；
  校对一变 hash 即变 → 重建。
- app.py /api/save 保存校对后调用 rebuild_units_for_image() 后台重建该图（M5）。

学术纪律（继承 reassembler）：
- LLM 输出字符多集校验（roster 允许补分隔标点）＋ line_ids 全覆盖校验；
  任一失败 → 单元降级 needs_review，绝不静默入库。
- needs_review 的单元不进 RAG、不进导出已验证视图（下游 phase 2/3 落实）。

CLI 用法：
    python unit_builder.py --xlsx output.xlsx                  # 全部图（存量重建）
    python unit_builder.py --image-stems <stem> ... --xlsx ...  # 指定图
    python unit_builder.py --project-id <pid> --no-llm          # 项目内图、纯规则
"""
from __future__ import annotations

import argparse
import hashlib
import json
import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple
log = logging.getLogger("unit_builder")

SCHEMA_VERSION = "1.0.0"

# 默认目录（与 exporter.py / rag.py / reassembler.py 对齐）
BASE_DIR = Path(__file__).parent.resolve()
DEFAULT_STRUCTURED_DIR = BASE_DIR / "data" / "structured"
DEFAULT_UNITS_DIR = BASE_DIR / "data" / "units"
DEFAULT_REASSEMBLED_DIR = BASE_DIR / "data" / "reassembled"
DEFAULT_XLSX = BASE_DIR / "output.xlsx"
DEFAULT_LLM_CONFIG = BASE_DIR / "data" / "llm_config.json"

# ---- 复用 reassembler 的既有资产（分类器 / 重排 / 校验）----
from reassembler import (  # noqa: E402
    ROSTER_FORMAT_CHARS,
    _text_of,
    _verify_coverage,
    box_of,
    classify_content_type,
    classify_layout,
    deterministic_reassemble,
    llm_classify_content_type,
    llm_reassemble_chunked,
    llm_reassemble_roster,
    verify_multiset,
)
import template_store  # M8b：用户样本模板（few-shot 注入 + get_active_template）

# ============================================
# 1. 中文数字金额解析（ledger 金额归一）
# ============================================
_CN_DIGIT = {"零": 0, "一": 1, "二": 2, "三": 3, "四": 4,
             "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
_CN_LOWUNIT = {"十": 10, "拾": 10, "百": 100, "千": 1000, "仟": 1000}
# 兩/两 之后的十进制小数位（錢=0.1 / 分=0.01 / 厘=0.001）
_CN_DEC = {"錢": 0.1, "钱": 0.1, "分": 0.01, "厘": 0.001, "毛": 0.001, "毫": 0.0001}

# 金额子串（用于"金额原文"提取）：连续的中文数字/单位字符，至少含一个銀衡单位
_AMOUNT_RUN_RE = re.compile(r"[一二三四五六七八九十百千萬万零兩两拾仟][一二三四五六七八九十百千萬万零兩两拾仟錢钱分厘毫]*")
_CURRENCY_RE = re.compile(r"(規元|規銀|库平|庫平|庫銀|库银|紋銀|纹银|洋銀|洋银|銀元|银元|銀两|银两|洋|銀|银|元)")
_DIR_RE = re.compile(r"[存該收欠支撥撥付領借還繳缴兑兌]")
_LEDGER_NAME_STOP = re.compile(r"[，,、。；;：:]")
def _cn_amount_to_float(text: str) -> Optional[float]:
    """中文金额串 → 兩 为单位的 float（"五十九萬七千三百十二兩九錢一分三厘"→597312.913）。

    单遍分级扫描：萬組内十/百/千进 section，遇 萬/万 封顶 ×10000 并入 total；
    「兩/两」是 multiplier=1 的收尾单位（3兩=3.0），但后跟位值字符时视为数字 2
    （"兩千"=2000）；錢/分/厘 按 0.1/0.01/0.001 累加。缺数字的小数位（OCR 掉字，
    如"錢分"）按 0 计，宁缺勿猜。解析不出任何数值 → None。
    """
    total = 0.0    # 已封頂的萬組 + 兩 后小数
    section = 0.0  # 当前萬組内的整数累计（千/百/十）
    cur: Optional[int] = None
    for i, ch in enumerate(text):
        if ch in _CN_DIGIT:
            cur = _CN_DIGIT[ch]
        elif ch in _CN_LOWUNIT:
            section += (cur if cur is not None else 1) * _CN_LOWUNIT[ch]
            cur = None
        elif ch in ("萬", "万"):
            section += cur or 0
            total += section * 10000
            section = 0.0
            cur = None
        elif ch in ("兩", "两"):
            nxt = text[i + 1] if i + 1 < len(text) else ""
            if nxt in _CN_LOWUNIT:
                cur = 2  # "兩千"：兩 = 数字 2
            else:
                total += section + (cur or 0)  # 兩 = multiplier 1 的收尾单位
                section = 0.0
                cur = None
        elif ch in _CN_DEC:
            if cur is not None:
                total += cur * _CN_DEC[ch]
                cur = None
            # 缺数字的小数位：按 0 计（不猜）
        # 零 / 未知字符（OCR 噪声）跳过
    return total if total > 0 else None


def extract_amount(text: str) -> Tuple[str, Optional[float]]:
    """从一行文本提取（金额原文, 金额归一）。无金额返回 ("", None)。"""
    best, best_len = "", 0
    for m in _AMOUNT_RUN_RE.finditer(text):
        run = m.group(0)
        if ("兩" in run or "两" in run or "錢" in run or "钱" in run) and len(run) > best_len:
            best, best_len = run, len(run)
    if not best:
        return "", None
    return best, _cn_amount_to_float(best)


# ============================================
# 2. ledger 判定与条目切分（规则，零 LLM）
# ============================================
def detect_ledger(lines: List[dict], corrected_map: Optional[Dict[int, str]] = None) -> bool:
    """规则判定是否账簿页：≥2 行含銀衡单位（兩/两）且 ≥2 行含收付方向字。

    判据保守（宁可漏判回 prose 路径，不可误判把散文切碎）；金额密度不足的
    混合页（正文里偶尔提到银两数）不会命中。
    """
    n_amount = n_dir = 0
    for ln in lines:
        t = _text_of(ln, corrected_map)
        if re.search(r"[兩两]", t) or re.search(r"[一二三四五六七八九十]{2,}[錢钱]分?", t):
            n_amount += 1
        if _DIR_RE.search(t):
            n_dir += 1
    return n_amount >= 2 and n_dir >= 2


def _ledger_reading_order(lines: List[dict]) -> List[dict]:
    """账簿阅读序：竖排栏右→左 = VL 扫描序（col_index 升序，栏内 row_index 升序）。"""
    return sorted(lines, key=lambda ln: (int(ln.get("col_index", 0)), int(ln.get("row_index", 0))))


def split_ledger_entries(lines: List[dict], corrected_map: Optional[Dict[int, str]] = None) -> List[dict]:
    """把账簿页的行切成账目条目（确定性状态机）。

    规则：含金额的行 = 账目的金额行（一条账目以金额行收尾）；金额行之前的
    连续非金额行 = 户名缓冲。状态机逐行推进，遇金额行即 flush 一条账目。

    Returns:
        [{"lines": [...], "has_amount": bool}]，每项一个候选条目（保持阅读序）
    """
    ordered = _ledger_reading_order(lines)
    groups: List[dict] = []
    buf: List[dict] = []  # 当前户名缓冲
    for ln in ordered:
        t = _text_of(ln, corrected_map)
        amt_raw, _ = extract_amount(t)
        if amt_raw:
            buf.append(ln)
            groups.append({"lines": buf, "has_amount": True})
            buf = []
        else:
            if groups and not buf and groups[-1]["has_amount"] and _looks_like_name(t):
                buf = [ln]  # 新条目从户名行开始
            else:
                buf.append(ln)
    if buf:
        groups.append({"lines": buf, "has_amount": False})
    return groups


def _looks_like_name(t: str) -> bool:
    """短行（2-8 字、无标点结尾、无金额）更像户名行的启发式。"""
    t = t.strip()
    return 2 <= len(t) <= 8 and not _LEDGER_NAME_STOP.search(t[-1])


def build_ledger_unit_fields(entry_lines: List[dict],
                             corrected_map: Optional[Dict[int, str]]) -> Tuple[dict, str]:
    """一条账目 → (fields, unit_text)。unit_text = 条目行文本依序拼接（不改字）。"""
    texts = [_text_of(ln, corrected_map) for ln in entry_lines]
    unit_text = "".join(texts)
    name_buf, amount_buf = [], []
    for ln, t in zip(entry_lines, texts):
        amt_raw, _ = extract_amount(t)
        (amount_buf if amt_raw else name_buf).append(t)
    dm = _DIR_RE.search(unit_text)
    cm = _CURRENCY_RE.search(unit_text)
    amount_raw, amount_norm = ("", None)
    for t in amount_buf:
        amount_raw, amount_norm = extract_amount(t)
        if amount_raw:
            break
    fields = {
        "户名": "".join(name_buf).strip(),
        "方向": dm.group(0) if dm else "",
        "币种": cm.group(0) if cm else "",
        "金额原文": amount_raw,
        "金额归一": amount_norm,
    }
    return fields, unit_text


# ============================================
# 3. roster 条目字段抽取（regex v1，不改字）
# ============================================
_ERA_STR = ("洪武|建文|永乐|洪熙|宣德|正統|正统|景泰|天順|天顺|成化|弘治|正德|嘉靖|隆慶|隆庆|"
            "萬曆|万历|泰昌|天啟|天启|崇禎|崇祯|順治|顺治|康熙|雍正|乾隆|嘉慶|嘉庆|道光|咸豐|咸丰|"
            "同治|光緒|光绪|宣統|宣统|民國|民国|同光|宣光")
_ERA_RE = re.compile(
    r"(?P<era>" + _ERA_STR + r")"
    r"(?P<rest>[甲乙丙丁戊己庚辛壬癸]{0,2}[子丑寅卯辰巳午未申酉戌亥]{0,2}"
    r"[初元零一二三四五六七八九十]{0,4})?[年科]?"
)
_GANZHI_RE = re.compile(r"[甲乙丙丁戊己庚辛壬癸][子丑寅卯辰巳午未申酉戌亥]")
_DEGREE_RE = re.compile(r"(進士|进士|舉人|举人|貢生|贡生|歲貢|岁贡|拔貢|拔贡|優貢|优贡|恩貢|恩贡|副貢|副贡|孝廉|狀元|状元|榜眼|探花|翰林|內閣|内阁|中書|中书)")
_NAME_STOP_RE = re.compile(r"[，,、。；;：:（(【\[]")
_SAME_AS_RE = re.compile(r"上同|同上")


def extract_roster_fields(text: str) -> dict:
    """名录条目文本 → 结构化字段（regex v1；抽不出的字段留空，宁缺勿猜）。"""
    fields: Dict[str, str] = {"姓名": "", "科年": "", "身份": "", "备注": ""}
    era_m = _ERA_RE.search(text)
    if era_m:
        fields["科年"] = era_m.group(0)
    gz_m = _GANZHI_RE.search(text)
    if gz_m and gz_m.group(0) not in fields["科年"]:
        fields["科年"] = (fields["科年"] + gz_m.group(0)) if fields["科年"] else gz_m.group(0)
    deg_m = _DEGREE_RE.search(text)
    if deg_m:
        fields["身份"] = deg_m.group(0)
    # 姓名：首个分隔符/科年起点之前的前缀（2-6 字才认；否则保守留空）
    head = text
    cut_positions = []
    if era_m:
        cut_positions.append(era_m.start())
    sep_m = _NAME_STOP_RE.search(text)
    if sep_m:
        cut_positions.append(sep_m.start())
    if cut_positions:
        head = text[: min(cut_positions)]
    head = head.strip()
    if 1 <= len(head) <= 6:
        fields["姓名"] = head
    # 备注：去掉已抽取部分后的剩余（保持原字符，不改字）
    remainder = text
    for taken in (fields["姓名"], fields["科年"], fields["身份"]):
        if taken:
            remainder = remainder.replace(taken, "", 1)
    fields["备注"] = remainder.strip("，,、。 ．.")
    return fields


# ============================================
# 4. 单元构建核心
# ============================================
def corrected_map_strict(xlsx_path: Path) -> Dict[int, str]:
    """读主工作簿 → {xlsx_id: corrected_text}，仅采纳"真校对"（corrected ≠ OCR 原文）。

    背景（Phase 1 实测，2026-08-31）：主工作簿为 VL 列模式单字/短段行，且
    corrected_text 列**预填 OCR 原文**——exporter.corrected_map_from_xlsx 会把
    7000+ 条"伪修正"（原文换了个行界而已）全读进来；unit_builder 按行套用会与
    L2 行粒度错位、污染字符多集校验。这里与 OCR 文本列做差集只保留真实编辑；
    残余的行粒度错位仍由 exporter.safe_corrected 粒度校验兜底（宁可漏接不错接）。
    """
    result: Dict[int, str] = {}
    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return result
        col_id = 0
        col_ocr = next((i for i, h in enumerate(header)
                        if h and str(h).strip() in ("OCR 文本", "ocr_text", "text")), 4)
        col_corr = next((i for i, h in enumerate(header)
                         if h and str(h).strip() in ("校对后文本", "corrected_text")), 7)
        for row in rows:
            rid, ocr, corr = row[col_id], row[col_ocr], row[col_corr]
            if not isinstance(rid, (int, float)) or not corr:
                continue
            corr_s = str(corr).strip()
            ocr_s = str(ocr or "").strip()
            if corr_s and corr_s != ocr_s:  # 真校对：与原文不同
                result[int(rid)] = corr_s
    except Exception as e:
        log.warning(f"[unit_builder] 读取主工作簿严格修正列失败（继续无校勘）: {e}")
    return result


def source_text_hash(lines: List[dict], corrected_map: Optional[Dict[int, str]] = None) -> str:
    """全图行文本（含校对合并后）sha1——缓存键：校对一变 hash 即变。"""
    joined = "\x1f".join(_text_of(ln, corrected_map) for ln in lines)
    return "sha1:" + hashlib.sha1(joined.encode("utf-8")).hexdigest()


def _unit_from_segment(seq: int, stem: str, seg: dict, lines_by_id: Dict[int, dict],
                       corrected_map: Optional[Dict[int, str]],
                       unit_type: str, fields: Optional[dict],
                       ignore_extra: Optional[set]) -> dict:
    """一个 segment（text + line_ids）→ unit dict（含逐字校验与溯源）。"""
    seg_lines = [lines_by_id[int(i)] for i in seg.get("line_ids", []) if int(i) in lines_by_id]
    input_texts = [_text_of(ln, corrected_map) for ln in seg_lines]
    vm = verify_multiset(input_texts, [seg.get("text", "")], ignore_extra=ignore_extra)
    xids = [int(ln.get("xlsx_id", 0)) for ln in seg_lines]
    boxes = [b for b in (box_of(ln) for ln in seg_lines) if b]
    box_union: Optional[List[float]] = None
    if boxes:
        box_union = [min(b[0] for b in boxes), min(b[1] for b in boxes),
                     max(b[2] for b in boxes), max(b[3] for b in boxes)]
    return {
        "unit_id": f"u_{stem}_{seq:04d}",
        "unit_type": unit_type,
        "text": seg.get("text", ""),
        "fields": fields or {},
        "line_ids": list(seg.get("line_ids", [])),
        "xlsx_id_range": [min(xids), max(xids)] if xids else [0, 0],
        "box": box_union,
        "verify": {
            "multiset": vm["matched"],
            "coverage": True,  # 页级覆盖校验统一做（见 build_units_image）
            "diff_total": vm["diff_total"],
            "status": "verified" if vm["matched"] else "needs_review",
        },
        "rebuilt_at": datetime.now().isoformat(timespec="seconds"),
    }


def _same_as_flag(unit: dict) -> None:
    """"上同/同上"指代标记（字段继承语义，phase 2 RAG 字段索引时展开）。"""
    if _SAME_AS_RE.search(unit.get("text", "")):
        unit["inherited"] = True


# ---- M8a：名录"同上"科年继承传播（确定性规则，用户定义语义 2026-09-01）----
# 用户语义："同上" = 该名字录取科年与纵列上一个人相同，依此上溯
# 直到直接列出录取年份的那个名字。实现：同列自上而下传播最近显式科年；
# 跨列不传播（宁缺勿猜原则——新列顶部若直接列年份会自然更新 carry）。
_COL_OVERLAP_TOL = 8.0  # px；x 区间重叠容差（OCR 框常有几像素抖动）


def _same_col(box_a: Optional[list], box_b: Optional[list]) -> bool:
    """两个 [x0,y0,x1,y1] box 是否属于同一纵列（x 区间有重叠）。"""
    if box_a is None or box_b is None:
        return True  # 坐标缺失时保守视为同列（退化为纯顺序继承）
    return max(box_a[0], box_b[0]) - _COL_OVERLAP_TOL <= min(box_a[2], box_b[2]) + _COL_OVERLAP_TOL


def propagate_roster_inheritance(units: List[dict]) -> int:
    """名录"同上"科年继承：就地补全 fields["科年"]，返回继承条数。

    - 只处理 roster_entry；文本不变（字符不动原则），只 enrich 字段
    - 同列（box x 区间重叠）内自上而下：显式科年更新 carry；"同上/上同"
      且科年缺失 → 继承 carry，并在 inherited 里记录溯源
    - 换列时 carry 重置；新列内"同上"无列内来源 → 保持科年为空（不跨列猜）
    - 继承结果仍标 needs_review 不变（不伪造 verified 资格）
    """
    carry: Optional[str] = None
    carry_box: Optional[list] = None
    n_inherited = 0
    for u in units:
        if u.get("unit_type") != "roster_entry":
            continue
        fields = u.setdefault("fields", {})
        box = u.get("box")
        text = u.get("text", "")
        if _SAME_AS_RE.search(text) and not fields.get("科年"):
            if carry and _same_col(box, carry_box):
                fields["科年"] = carry
                u["inherited"] = {"field": "科年", "via": "同上", "value": carry}
                n_inherited += 1
                carry_box = box or carry_box  # 继承条目也算列内锚点
        elif fields.get("科年"):
            carry = fields["科年"]
            carry_box = box
    return n_inherited


def reuse_l4_segments(
    stem: str,
    l2: List[dict],
    corrected_map: Optional[Dict[int, str]] = None,
    reassembled_dir: Path = DEFAULT_REASSEMBLED_DIR,
    ignore_extra: Optional[set] = None,
) -> Optional[List[dict]]:
    """复用 reassembler 的已验证 L4 产物（data/reassembled/<stem>.json）。

    这是方案 M1 的关键设计：reassembler 的 LLM 条目化结果（实测成功率 ~1/5、
    单图可耗时 169s+）是最贵的资产，units 层直接继承而非重算——消解 G5 成本。

    采纳门槛（不满足则返回 None，调用方走 LLM/规则重算）：
    - L4 method == "llm"（规则重排不值得复用，重算零成本）
    - 对**当前**行文本（含最新校对）重跑字符多集校验 + line_ids 覆盖校验，
      全部通过——校对一变 L4 即失效，保证不把陈旧重排结果当最新事实
    """
    path = Path(reassembled_dir) / f"{stem}.json"
    if not path.exists():
        return None
    try:
        l4 = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as e:
        log.warning(f"[unit_builder] {stem}: L4 读取失败 ({e})")
        return None
    if l4.get("method") != "llm":
        return None
    segs = l4.get("segments") or []
    if not segs:
        return None
    input_texts = [_text_of(ln, corrected_map) for ln in l2]
    vm = verify_multiset(input_texts, [s.get("text", "") for s in segs], ignore_extra=ignore_extra)
    cov_ok, missing, extra = _verify_coverage(segs, l2)
    if not (vm["matched"] and cov_ok):
        log.info(f"[unit_builder] {stem}: L4 与当前文本不匹配（multiset={vm['matched']} "
                 f"coverage={cov_ok}），走重算")
        return None
    log.info(f"[unit_builder] {stem}: 复用已验证 L4 重排结果（{len(segs)} 段）")
    return segs


def _page_coverage(units: List[dict], lines: List[dict]) -> Tuple[bool, List[int], List[int]]:
    """页级 line_ids 覆盖校验：所有单元的 line_ids 恰好覆盖全部行一次。"""
    used: set = set()
    dup: set = set()
    for u in units:
        for i in u["line_ids"]:
            if i in used:
                dup.add(i)
            used.add(i)
    expected = {int(ln.get("line_id")) for ln in lines if ln.get("line_id") is not None}
    missing = sorted(expected - used)
    extra = sorted(used - expected)
    return (not missing and not extra and not dup), missing + sorted(dup), extra


def build_units_image(
    stem: str,
    structured_dir: Path = DEFAULT_STRUCTURED_DIR,
    corrected_map: Optional[Dict[int, str]] = None,
    client=None,
    out_dir: Optional[Path] = DEFAULT_UNITS_DIR,
    force: bool = False,
    roster_max_lines: int = 30,
    reassembled_dir: Optional[Path] = DEFAULT_REASSEMBLED_DIR,
) -> Optional[dict]:
    """构建单张图的语义单元并落盘 data/units/<stem>.json。

    Args:
        stem: 图片 stem（structured 文件名）
        corrected_map: {xlsx_id: corrected_text}（人工校对；空则用 OCR 原文）
        client: LLM client（None = 纯规则；roster 无 LLM 时降级 needs_review）
        force: True 跳过 hash 缓存强制重算
        out_dir: None = 不落盘（纯计算）
        reassembled_dir: L4 复用目录；None = 禁用复用

    Returns:
        units payload dict；structured 缺失返回 None
    """
    corrected_map = corrected_map or {}
    path = structured_dir / f"{stem}.json"
    if not path.exists():
        log.warning(f"[unit_builder] structured 不存在: {path}")
        return None
    data = json.loads(path.read_text(encoding="utf-8"))
    l0 = data.get("L0_document", {}) or {}
    l2 = data.get("L2_lines", []) or []
    if not l2:
        log.warning(f"[unit_builder] 无 L2 行，跳过: {stem}")
        return None

    s_hash = source_text_hash(l2, corrected_map)
    out_path = (out_dir / f"{stem}.json") if out_dir else None
    if out_path and out_path.exists() and not force:
        try:
            cached = json.loads(out_path.read_text(encoding="utf-8"))
            if cached.get("source_text_hash") == s_hash:
                log.info(f"[unit_builder] {stem}: hash 命中缓存，跳过")
                cached["_cached"] = True  # 仅内存标记，不落盘
                return cached
        except (json.JSONDecodeError, OSError) as e:
            log.warning(f"[unit_builder] {stem}: 缓存读取失败，重算 ({e})")

    layout = classify_layout(l2)
    image_name = l0.get("image_name", f"{stem}.png")
    lines_by_id: Dict[int, dict] = {int(ln.get("line_id")): ln for ln in l2 if ln.get("line_id") is not None}

    # ---- 内容类型路由 ----
    # ledger 优先（规则，零成本）：防止把账簿页误判成 prose/ambiguous 后走重排路径
    if detect_ledger(l2, corrected_map):
        content_type, classify_method, classify_reason = "ledger", "rule", "銀衡单位+收付方向字命中"
    else:
        content_type = classify_content_type(l2)
        classify_method, classify_reason = "rule", ""
        if client is not None and content_type != "roster":
            # roster 规则判据实测标定可靠，无需 LLM 复核；其余类型 LLM 裁决
            try:
                llm_ct, reason = llm_classify_content_type(client, l2, image_name)
                if llm_ct != "ambiguous":
                    content_type, classify_method, classify_reason = llm_ct, "llm", reason
                else:
                    classify_reason = f"LLM 无法判定（{reason}），回落规则: {content_type}"
            except Exception as e:
                log.warning(f"[unit_builder] {stem}: LLM 分类失败，用规则结果 ({e})")

    units: List[dict] = []
    build_method = "rule"

    if content_type == "ledger":
        # ---- 账目条目（规则状态机，字符不动 → 校验天然通过）----
        for g in split_ledger_entries(l2, corrected_map):
            entry_lines = [ln for ln in g["lines"] if int(ln.get("line_id", -1)) in lines_by_id]
            if not entry_lines:
                continue
            if g["has_amount"]:
                fields, unit_text = build_ledger_unit_fields(entry_lines, corrected_map)
                utype = "ledger_entry"
            else:
                fields, unit_text = {}, "".join(_text_of(ln, corrected_map) for ln in entry_lines)
                utype = "text_fragment"
            seg = {"text": unit_text, "line_ids": [int(ln.get("line_id")) for ln in entry_lines]}
            u = _unit_from_segment(len(units) + 1, stem, seg, lines_by_id, corrected_map, utype, fields, None)
            units.append(u)
    elif content_type == "roster":
        # ---- 名录条目：优先复用已验证 L4 → LLM 聚合 → 规则降级 needs_review ----
        segs = (reuse_l4_segments(stem, l2, corrected_map, reassembled_dir,
                                  ignore_extra=ROSTER_FORMAT_CHARS)
                if reassembled_dir else None)
        if segs is not None:
            build_method = "llm_reused"
        elif client is not None:
            try:
                # M8b：激活的 roster 样本模板 → few-shot 注入
                tmpl = template_store.get_active_template("roster")
                segs = llm_reassemble_roster(client, l2, corrected_map, image_name,
                                             max_lines=roster_max_lines, template=tmpl)
                build_method = "llm"
            except Exception as e:
                log.warning(f"[unit_builder] {stem}: 名录重组失败，规则降级 ({e})")
                segs = None
        if segs is None:
            segs = deterministic_reassemble(l2, corrected_map,
                                            layout if layout in ("vertical", "horizontal") else "vertical")
            for s in segs:
                s["status"] = "needs_review"  # 规则分栏 ≠ 条目化，统一降级
        for i, s in enumerate(segs, start=1):
            u = _unit_from_segment(i, stem, s, lines_by_id, corrected_map, "roster_entry",
                                   extract_roster_fields(s.get("text", "")),
                                   ignore_extra=ROSTER_FORMAT_CHARS)
            if s.get("status") == "needs_review":
                u["verify"]["status"] = "needs_review"
                u["verify"]["reason"] = "no_llm_rule_fallback" if client is None else "llm_failed_fallback"
            _same_as_flag(u)
            units.append(u)
        # M8a："同上"科年继承传播（确定性规则；文本不动、只补字段）
        n_inherited = propagate_roster_inheritance(units)
        if n_inherited:
            log.info(f"[unit_builder] {stem}: 同上科年继承 {n_inherited} 条")
    else:
        # ---- prose / letter / document / table / mixed / ambiguous → 段落 ----
        seg_layout = layout if layout in ("vertical", "horizontal") else "ambiguous"
        segs = None
        if seg_layout == "ambiguous":
            # 歧义版面的 LLM 重排同样昂贵，优先复用已验证 L4
            segs = (reuse_l4_segments(stem, l2, corrected_map, reassembled_dir)
                    if reassembled_dir else None)
            if segs is not None:
                build_method = "llm_reused"
        if segs is None and seg_layout == "ambiguous" and client is not None:
            try:
                segs, _ = llm_reassemble_chunked(client, l2, corrected_map, seg_layout, image_name)
                build_method = "llm"
            except Exception as e:
                log.warning(f"[unit_builder] {stem}: LLM 重排失败，规则回退 ({e})")
                segs = None
        if segs is None:
            segs = deterministic_reassemble(l2, corrected_map, seg_layout)
        for i, s in enumerate(segs, start=1):
            u = _unit_from_segment(i, stem, s, lines_by_id, corrected_map, "paragraph", None, None)
            units.append(u)

    # ---- 页级覆盖校验：任一失败全部单元降级 ----
    cov_ok, missing, extra = _page_coverage(units, l2)
    review_flags: List[str] = []
    if not cov_ok:
        review_flags.append(f"line_ids 覆盖不完整: missing/dup={missing[:8]} extra={extra[:8]}")
        for u in units:
            u["verify"]["coverage"] = False
            u["verify"]["status"] = "needs_review"
    n_bad = sum(1 for u in units if u["verify"]["status"] != "verified")

    payload = {
        "_schema_version": SCHEMA_VERSION,
        "image_stem": stem,
        "image_name": image_name,
        "source_text_hash": s_hash,
        "layout": layout,
        "content_type": content_type,
        "classify_method": classify_method,
        "classify_reason": classify_reason,
        "build_method": build_method,
        "unit_count": len(units),
        "verified_count": len(units) - n_bad,
        "needs_review_count": n_bad,
        "built_at": datetime.now().isoformat(timespec="seconds"),
        "units": units,
    }
    if review_flags:
        payload["review_flags"] = review_flags

    if out_path:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        tmp = out_path.with_suffix(".json.tmp")
        tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        tmp.replace(out_path)
    log.info(f"[unit_builder] {stem}: type={content_type} method={build_method} "
             f"units={len(units)} verified={len(units) - n_bad}")
    return payload


# ============================================
# 5. 校对回流入口（app.py /api/save 调用）
# ============================================
def rebuild_units_for_image(
    image_name: str,
    xlsx_path: Path = DEFAULT_XLSX,
    structured_dir: Path = DEFAULT_STRUCTURED_DIR,
    out_dir: Path = DEFAULT_UNITS_DIR,
    use_llm: bool = True,
    force: bool = True,
    llm_config_path: Path = DEFAULT_LLM_CONFIG,
) -> Optional[dict]:
    """单图重建（校对保存后的后台入口）：取 xlsx 校正 → 重建 units。

    - use_llm=True 时尝试加载 LLM client（roster 重组需要）；失败降级纯规则
    - 纯规则且该图已有 LLM 构建的 roster units 时跳过（避免用劣化结果覆盖好结果）
    - 异常只记日志，绝不上抛（不阻塞 /api/save）

    ✅ P-K9 阶段 0c（2026-09-10）：新增 ``force``。
    - ``force=True``（默认，/api/save 校对回流语义）：**忽略 hash 缓存**强制重算，
      因为保存动作本身就是"内容已变"的信号，且待重建的图量小。
    - ``force=False``（自动构建语义）：走 ``build_units_image`` 的
      ``source_text_hash`` 缓存——未校对且已构建过的图**零成本秒回**，
      这是"OCR 完成后自动构建"能安全常驻的前提。
    """
    stem = Path(image_name).stem
    try:
        corrected: Dict[int, str] = {}
        if xlsx_path and Path(xlsx_path).exists():
            corrected = corrected_map_strict(Path(xlsx_path))
        client = None
        if use_llm:
            try:
                from llm_client import get_or_create_client, load_config
                client = get_or_create_client(load_config(Path(llm_config_path)))
            except Exception as e:
                log.warning(f"[unit_builder] {stem}: LLM 不可用，规则路径 ({e})")
        if client is None:
            existing = out_dir / f"{stem}.json"
            if existing.exists():
                try:
                    old = json.loads(existing.read_text(encoding="utf-8"))
                    if old.get("build_method") == "llm" and old.get("content_type") == "roster":
                        log.info(f"[unit_builder] {stem}: 无 LLM，跳过以免覆盖 LLM 名录条目")
                        return None
                except (json.JSONDecodeError, OSError):
                    pass
        return build_units_image(stem, structured_dir, corrected, client=client,
                                 out_dir=out_dir, force=force)
    except Exception as e:
        log.warning(f"[unit_builder] {stem}: 重建失败（不影响校对保存）: {e}")
        return None


# ============================================
# 6. 项目/全量入口 + CLI
# ============================================
def _project_image_stems(project_id: str, data_dir: Path = BASE_DIR / "data") -> List[str]:
    """从 data/project_assignments.json 取项目图片白名单（按 stem）。

    ✅ M23：核心逻辑委托 data_io.get_project_image_stems 统一实现（消三胞胎重复）。
    """
    from data_io import get_project_image_stems
    return get_project_image_stems(project_id, data_dir)


def update_index(units_dir: Path, payloads: List[dict]) -> None:
    """重建/更新 data/units/_index.json（stem 级摘要，供 dashboard/导出用）。"""
    idx_path = units_dir / "_index.json"
    index: Dict[str, dict] = {}
    if idx_path.exists():
        try:
            old = json.loads(idx_path.read_text(encoding="utf-8"))
            for it in old.get("images", []):
                index[it["image_stem"]] = it
        except (json.JSONDecodeError, OSError, KeyError):
            index = {}
    for p in payloads:
        if not p:
            continue
        index[p["image_stem"]] = {
            "image_stem": p["image_stem"],
            "image_name": p["image_name"],
            "content_type": p["content_type"],
            "build_method": p["build_method"],
            "unit_count": p["unit_count"],
            "verified_count": p["verified_count"],
            "needs_review_count": p["needs_review_count"],
            "source_text_hash": p["source_text_hash"],
            "built_at": p["built_at"],
        }
    payload = {
        "_schema_version": SCHEMA_VERSION,
        "updated_at": datetime.now().isoformat(timespec="seconds"),
        "image_count": len(index),
        "images": sorted(index.values(), key=lambda it: it["image_stem"]),
    }
    tmp = idx_path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(idx_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="语义单元层构建器（unit layer / L3）")
    parser.add_argument("--project-id", default="", help="项目 ID（取该项目图片白名单）")
    parser.add_argument("--image-stems", nargs="*", default=None,
                        help="图片白名单（优先于 --project-id；缺省=structured 全部）")
    parser.add_argument("--xlsx", default=str(DEFAULT_XLSX), help="主工作簿（人工校对优先）")
    parser.add_argument("--structured-dir", default=str(DEFAULT_STRUCTURED_DIR))
    parser.add_argument("--out-dir", default=str(DEFAULT_UNITS_DIR))
    parser.add_argument("--llm-config", default=str(DEFAULT_LLM_CONFIG))
    parser.add_argument("--no-llm", action="store_true", help="纯规则（roster 降级 needs_review）")
    parser.add_argument("--force", action="store_true", help="忽略 hash 缓存强制重算")
    parser.add_argument("--roster-max-lines", type=int, default=30,
                        help="名录条目化分块行数（同 reassembler，默认 30）")
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(message)s")
    structured_dir = Path(args.structured_dir)
    if not structured_dir.is_dir():
        log.error(f"structured 目录不存在: {structured_dir}")
        return 2

    stems = args.image_stems
    if stems is None and args.project_id:
        stems = _project_image_stems(args.project_id)
        log.info(f"项目 {args.project_id} 图片白名单: {len(stems)} 张")
    elif stems is None:
        stems = sorted(p.stem for p in structured_dir.glob("*.json") if not p.name.startswith("_"))
    if not stems:
        log.warning("无图可构建")
        return 1

    corrected: Dict[int, str] = {}
    xlsx_path = Path(args.xlsx)
    if xlsx_path.exists():
        corrected = corrected_map_strict(xlsx_path)
        log.info(f"主工作簿严格修正列（真校对）: {len(corrected)} 条")

    client = None
    if not args.no_llm:
        try:
            from llm_client import get_or_create_client, load_config
            client = get_or_create_client(load_config(Path(args.llm_config)))
            log.info(f"LLM 可用: {client.model_name}")
        except Exception as e:
            log.warning(f"LLM 不可用（{e}）——roster 将降级 needs_review")

    out_dir = Path(args.out_dir)
    payloads: List[dict] = []
    n_cache = n_ok = n_fail = 0
    for stem in stems:
        try:
            p = build_units_image(stem, structured_dir, corrected, client=client,
                                  out_dir=out_dir, force=args.force,
                                  roster_max_lines=args.roster_max_lines)
        except Exception as e:
            log.warning(f"[unit_builder] {stem}: 构建异常 {e}")
            p = None
        if p is None:
            n_fail += 1
        else:
            payloads.append(p)
            if p.get("_cached"):
                n_cache += 1
            else:
                n_ok += 1
    if payloads:
        update_index(out_dir, payloads)
    log.info(f"完成：共 {len(stems)} 张，新建/重算={n_ok}，失败/跳过={n_fail} → {out_dir}")
    return 0 if n_fail == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
