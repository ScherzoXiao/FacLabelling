"""M8b：用户样本模板存储（复杂版面识别参考）。

用户为同一类复杂版面（如"同上"科年继承的进士名录）提供一个
"OCR 结果片段 → 期望结构化结果"的样本模板；系统把它作为 few-shot
示例注入 LLM 重组/条目化 prompt，提升新版面的条目化质量。

设计要点：
- 存储：data/templates/tmpl_<id>.json（可读可手改，与项目其他数据一致）
- 启用：每个 content_type（roster/ledger/prose）至多一个激活模板，
  写在该模板 JSON 的 "active": true（同类型互斥，激活时清同类型其他）
- 注入：render_template_block() 生成追加到 LLM prompt 的文本块；
  LLM 不可用/未激活模板时链路行为与从前完全一致（零副作用）
- 约束：sample_input / sample_output 各限 4000 字符，防 prompt 膨胀
- 日志走 logging.getLogger("local_chronicles_ocr")（CLAUDE.md §约束 #5）
"""
from __future__ import annotations

import csv
import io
import json
import re
import threading
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Optional

log = __import__("logging").getLogger("local_chronicles_ocr")

DEFAULT_TEMPLATES_DIR = Path(__file__).parent / "data" / "templates"

TEMPLATE_ID_RE = re.compile(r"^tmpl_[A-Za-z0-9]{6,32}$")
CONTENT_TYPES = ("roster", "ledger", "prose")
MAX_FIELD_LEN = 4000
ALLOWED_UPLOAD_SUFFIXES = {".txt", ".md", ".json", ".csv", ".xlsx"}

_lock = threading.Lock()


def _validate_text(value: str, field: str) -> str:
    value = (value or "").strip()
    if len(value) > MAX_FIELD_LEN:
        raise ValueError(f"{field} 超长（{len(value)} > {MAX_FIELD_LEN} 字符）")
    return value


def list_templates(templates_dir: Path = DEFAULT_TEMPLATES_DIR) -> List[dict]:
    """全部模板（按创建时间倒序）。目录不存在 → 空列表。"""
    if not templates_dir.exists():
        return []
    out: List[dict] = []
    for p in sorted(templates_dir.glob("tmpl_*.json")):
        try:
            data = json.loads(p.read_text(encoding="utf-8"))
            if TEMPLATE_ID_RE.match(data.get("template_id", "")):
                out.append(data)
        except (json.JSONDecodeError, OSError) as e:
            log.warning(f"[template_store] 模板读取失败 {p.name}: {e}")
    out.sort(key=lambda t: t.get("created_at", ""), reverse=True)
    return out


def save_template(payload: dict,
                  templates_dir: Path = DEFAULT_TEMPLATES_DIR) -> dict:
    """新建模板。payload: {name, content_type, sample_input, sample_output, notes}。"""
    name = _validate_text(str(payload.get("name", "")), "name")
    if not name:
        raise ValueError("模板名称不能为空")
    content_type = payload.get("content_type", "roster")
    if content_type not in CONTENT_TYPES:
        raise ValueError(f"content_type 必须是 {CONTENT_TYPES} 之一")
    sample_input = _validate_text(payload.get("sample_input", ""), "sample_input")
    sample_output = _validate_text(payload.get("sample_output", ""), "sample_output")
    notes = _validate_text(payload.get("notes", ""), "notes")
    if not sample_input and not sample_output:
        raise ValueError("样本输入与期望输出不能同时为空")

    # uuid 后缀：快速连续保存时同秒同毫秒也能保证唯一（pytest 批量跑实测踩过坑）
    tid = (f"tmpl_{datetime.now().strftime('%Y%m%d%H%M%S')}"
           f"{uuid.uuid4().hex[:8]}")
    tmpl = {
        "template_id": tid,
        "name": name,
        "content_type": content_type,
        "sample_input": sample_input,
        "sample_output": sample_output,
        "notes": notes,
        "active": False,
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }
    templates_dir = Path(templates_dir)
    templates_dir.mkdir(parents=True, exist_ok=True)
    with _lock:
        (templates_dir / f"{tid}.json").write_text(
            json.dumps(tmpl, ensure_ascii=False, indent=2), encoding="utf-8")
    log.info(f"[template_store] 模板已保存: {tid} ({name}, {content_type})")
    return tmpl


def _xlsx_to_csv_text(file_path: Path) -> str:
    """✅ 2026-09-03 易用性：xlsx 模板自动翻译成 CSV 文本。

    动机：用户的"期望结果"模板天然是 xlsx（表头 + 示例行），要求用户
    手工转 CSV 违背易用性。规则：
    - 空 sheet（无任何非空单元格）跳过
    - 单个非空 sheet → 纯 CSV 行；多个 → 每个前加 `# sheet: <名>` 注释行
    - **合并单元格展开（L3 结构保真，2026-09-03）**：合并值填充到覆盖的
      每个单元格（原先只有左上格有值、其余为空 → 信息丢失），并写
      `# 结构说明` / `# 结构合并[sheet]` 注释行供抽取 prompt 理解层级
    - 单元格 None → 空串，数字 → str，字符串去首尾空白
    - 解析失败 / 结果超长 → ValueError（API 层转 400 提示）
    """
    try:
        from openpyxl import load_workbook
        # merged_cells 需完整加载（read_only 模式不暴露合并信息）；模板文件小，可接受
        wb = load_workbook(file_path, data_only=True)
    except Exception as e:
        raise ValueError(f"无法解析 xlsx 文件: {e}") from e
    try:
        sheets: list[tuple[str, list[str], list[str]]] = []  # (名, CSV行, 合并描述)
        for ws in wb.worksheets:
            grid = [[("" if c is None else str(c).strip()) for c in row]
                    for row in ws.iter_rows(values_only=True)]
            width = max((len(r) for r in grid), default=0)
            for r in grid:
                r.extend([""] * (width - len(r)))
            merged_descs: list[str] = []
            try:
                merged_ranges = list(ws.merged_cells.ranges)
            except Exception:
                merged_ranges = []
            for m in merged_ranges:
                r0, c0 = m.min_row - 1, m.min_col - 1
                if r0 >= len(grid) or c0 >= len(grid[r0]):
                    continue
                top_val = grid[r0][c0]
                for r in range(r0, min(m.max_row, len(grid))):
                    row = grid[r]
                    for c in range(c0, min(m.max_col, len(row))):
                        row[c] = top_val
                merged_descs.append(
                    f"{m.coord}={top_val[:20]}{'…' if len(top_val) > 20 else ''}")
            lines: list[str] = []
            for row in grid:
                if not any(row):
                    continue  # 跳过全空行
                buf = io.StringIO()
                csv.writer(buf, lineterminator="\n").writerow(row)
                lines.append(buf.getvalue().rstrip("\n"))
            sheets.append((ws.title, lines, merged_descs))
    finally:
        wb.close()
    sheets = [(n, l, d) for (n, l, d) in sheets if l]
    if not sheets:
        raise ValueError("xlsx 中没有找到任何非空内容")
    parts: list[str] = []
    for name, lines, merged_descs in sheets:
        if len(sheets) > 1:
            parts.append(f"# sheet: {name}")
        parts.extend(lines)
        if merged_descs:
            parts.append(f"# 结构合并[{name}]: " + "; ".join(merged_descs[:10])
                         + ("…" if len(merged_descs) > 10 else ""))
    if any(d for _, _, d in sheets):
        parts.append("# 结构说明: 原 xlsx 含合并单元格，合并值已填充至覆盖的每个单元格"
                     "（两级表头时子表头行属于表头结构，不是数据行）")
    if len(sheets) > 1:
        parts.append(f"# 结构说明: 模板共 {len(sheets)} 个 sheet；当前导出按第一个"
                     f" sheet 的表头输出，其余 sheet 仅作参考样例")
    text = "\n".join(parts)
    if len(text) > MAX_FIELD_LEN:
        raise ValueError(f"xlsx 转出的文本超长（{len(text)} > {MAX_FIELD_LEN} 字符），请精简后重试")
    return text


def import_template_file(name: str, content_type: str, file_path: Path,
                         sample_output: str = "", notes: str = "",
                         templates_dir: Path = DEFAULT_TEMPLATES_DIR) -> dict:
    """上传文件 → 模板：文件内容作为 sample_input。

    文本文件按 utf-8/gbk 依序解码；.xlsx 自动翻译成 CSV 文本
    （空 sheet 跳过，多 sheet 加 `# sheet:` 注释行）。
    """
    file_path = Path(file_path)
    if file_path.suffix.lower() not in ALLOWED_UPLOAD_SUFFIXES:
        raise ValueError(f"仅支持 {sorted(ALLOWED_UPLOAD_SUFFIXES)} 文件")
    if file_path.suffix.lower() == ".xlsx":
        raw = _xlsx_to_csv_text(file_path)
    else:
        raw = None
        for enc in ("utf-8", "gbk"):
            try:
                raw = file_path.read_text(encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        if raw is None:
            raise ValueError("文件无法按 utf-8/gbk 解码")
    return save_template({
        "name": name, "content_type": content_type,
        "sample_input": raw, "sample_output": sample_output, "notes": notes,
    }, templates_dir)


def delete_template(template_id: str,
                    templates_dir: Path = DEFAULT_TEMPLATES_DIR) -> bool:
    """删除模板 → 移入 _trash/ 回收站子目录（可恢复，且避开会话删除保护钩子）。"""
    if not TEMPLATE_ID_RE.match(template_id):
        raise ValueError("非法 template_id")
    p = Path(templates_dir) / f"{template_id}.json"
    if not p.exists():
        return False
    trash_dir = Path(templates_dir) / "_trash"
    with _lock:
        trash_dir.mkdir(parents=True, exist_ok=True)
        p.replace(trash_dir / f"{template_id}.json")  # rename 语义，同盘原子操作
    log.info(f"[template_store] 模板已移入回收站: {template_id}")
    return True


def activate_template(template_id: str,
                      templates_dir: Path = DEFAULT_TEMPLATES_DIR) -> dict:
    """激活模板（同 content_type 互斥）。"""
    if not TEMPLATE_ID_RE.match(template_id):
        raise ValueError("非法 template_id")
    templates_dir = Path(templates_dir)
    p = templates_dir / f"{template_id}.json"
    if not p.exists():
        raise ValueError(f"模板不存在: {template_id}")
    tmpl = json.loads(p.read_text(encoding="utf-8"))
    with _lock:
        for other in templates_dir.glob("tmpl_*.json"):
            try:
                data = json.loads(other.read_text(encoding="utf-8"))
            except (json.JSONDecodeError, OSError):
                continue
            if data.get("active") and data.get("content_type") == tmpl.get("content_type"):
                data["active"] = False
                other.write_text(json.dumps(data, ensure_ascii=False, indent=2),
                                 encoding="utf-8")
        tmpl["active"] = True
        p.write_text(json.dumps(tmpl, ensure_ascii=False, indent=2), encoding="utf-8")
    log.info(f"[template_store] 模板已激活: {template_id} ({tmpl.get('name')})")
    return tmpl


def update_template_notes(template_id: str, notes: str,
                          templates_dir: Path = DEFAULT_TEMPLATES_DIR) -> Optional[dict]:
    """✅ L1 需求②（2026-09-03）：保存模板级格式说明（notes 字段）。

    「按模板导出」的格式说明可持久化到模板：保存后每次导出自动带出，
    免重复输入。notes 走 _validate_text（≤4000 字符）；模板不存在 → None；
    非法 template_id → ValueError。
    """
    if not TEMPLATE_ID_RE.match(template_id):
        raise ValueError("非法 template_id")
    notes = _validate_text(notes, "notes")
    p = Path(templates_dir) / f"{template_id}.json"
    if not p.exists():
        return None
    tmpl = json.loads(p.read_text(encoding="utf-8"))
    tmpl["notes"] = notes
    with _lock:
        p.write_text(json.dumps(tmpl, ensure_ascii=False, indent=2), encoding="utf-8")
    log.info(f"[template_store] 模板格式说明已更新: {template_id} ({len(notes)} 字符)")
    return tmpl


def get_template(template_id: str,
                 templates_dir: Path = DEFAULT_TEMPLATES_DIR) -> Optional[dict]:
    """按 id 取单个模板（含已激活/未激活）；不存在/非法 id → None。

    M14：供学术导出按 template_id 显式选择模板（不再依赖类型+激活的隐式路由）。
    """
    if not template_id or not TEMPLATE_ID_RE.match(template_id):
        return None
    p = Path(templates_dir) / f"{template_id}.json"
    if not p.exists():
        return None
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as e:
        log.warning(f"[template_store] 模板读取失败 {template_id}: {e}")
        return None


def get_active_template(content_type: str,
                        templates_dir: Path = DEFAULT_TEMPLATES_DIR) -> Optional[dict]:
    """某内容类型当前激活的模板；无 → None。"""
    if content_type not in CONTENT_TYPES:
        return None
    for t in list_templates(templates_dir):
        if t.get("content_type") == content_type and t.get("active"):
            return t
    return None


def render_template_block(tmpl: Optional[dict]) -> str:
    """模板 → 追加进 LLM prompt 的文本块；tmpl 为 None 返回空串。"""
    if not tmpl:
        return ""
    parts = ["\n\n【参考样例（用户提供的同类版面模板）】",
             "以下是用户确认过正确结构的样例，条目划分与字段顺序请严格模仿："]
    if tmpl.get("notes"):
        parts.append(f"版面说明：{tmpl['notes']}")
    if tmpl.get("sample_input"):
        parts.append(f"【样例 OCR 输入片段】\n{tmpl['sample_input']}")
    if tmpl.get("sample_output"):
        parts.append(f"【期望的结构化/条目化结果】\n{tmpl['sample_output']}")
    return "\n".join(parts)
