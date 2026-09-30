# -*- coding: utf-8 -*-
"""栏目信息汇总模板解析（2026-09-06 工作流归位）。

动机：预期工作流中，用户在 OCR 结果出来后即提供「信息汇总模板」
（.docx / .xlsx / .md），系统据其表头提取手动标注可能用到的属性。
该入口原先误置于成果阶段的知识库（学术导出「提供新模板」），现迁到
专项栏目——本模块只负责「文件 → 文本 → 表头属性名」这一段纯解析；
存储走 data_io.update_project，档案 upsert 由 app.py 编排 profile_store。

2026-09-12 扩展：**一个栏目可以有多套模板样式**（= 多套属性集 = 多份方案）。
理由是执行层的"可拆卸"意义正在于「一套材料配多套规则」：同一批材料，
第一套规则强调公司名、第二套强调注册人。故：
  · 一份"样式"= `{label, filename, attrs, text, ts, profile_id}`
    —— `label` = 模板文件名去扩展名，是样式在栏目内的**身份**（同 label 视为同一套）。
  · 列表存 `project["summary_templates"]`；旧字段 `summary_template`（单值）
    **保留为"最近一次"**，供学术导出等既有消费口使用（向后兼容，零迁移）。
  · 样式 ↔ 档案：同 label 的样式并入**同一个**档案（即"改一改再传"），
    新 label 建**独立**档案（即"另一套规则"）。
本模块只做**纯数据变换**（不读盘、不写盘），谁存谁编排由 app.py 负责。
"""
from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Dict, List, Optional

ALLOWED_SUFFIXES = {".docx", ".xlsx", ".md"}
MAX_TEXT_LEN = 20000      # 提取文本上限（字符），防超大文件拖慢请求
DEFAULT_MAX_ATTRS = 64    # 与 profile_store.MAX_ATTRS 对齐（延迟导入前先本地兜底）

STYLES_FIELD = "summary_templates"   # 列表字段（新）
LEGACY_FIELD = "summary_template"    # 单值字段（旧，保留为"最近一次"）
LABEL_SEP = "｜"                     # 档案名里的栏目/样式分隔（全角竖线，与标点习惯一致）
_UNNAMED = "未命名样式"


def _read_text_file(file_path: Path) -> str:
    """文本文件按 utf-8/gbk 依序解码（与 template_store 同策略）。"""
    raw = None
    for enc in ("utf-8", "gbk"):
        try:
            raw = file_path.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        raise ValueError("文件无法按 utf-8/gbk 解码")
    return raw


def extract_text(file_path) -> str:
    """上传文件 → 纯文本。

    - .md：直接读文本
    - .xlsx：复用 template_store 的 CSV 转换（合并单元格已填充、
      多 sheet 加 `# sheet:` 注释行——正好被表头提取的注释跳过规则忽略）
    - .docx：python-docx 取**首个表格的首行**（合并单元格产生的重复
      单元格由表头去重兜住）；无表格 → 明确报错引导用户改用表格
    """
    file_path = Path(file_path)
    suffix = file_path.suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise ValueError(f"仅支持 {sorted(ALLOWED_SUFFIXES)} 文件")

    if suffix == ".xlsx":
        import template_store  # 延迟导入避免循环依赖
        text = template_store._xlsx_to_csv_text(file_path)
    elif suffix == ".md":
        text = _read_text_file(file_path)
    else:  # .docx
        try:
            import docx as _docx  # python-docx
        except ImportError as e:
            raise ValueError("服务器缺少 python-docx，无法解析 .docx 文件") from e
        try:
            doc = _docx.Document(str(file_path))
        except Exception as e:
            raise ValueError(f"无法解析 docx 文件: {e}") from e
        if not doc.tables:
            raise ValueError(
                "docx 中没有找到表格——请在模板中用表格列出信息属性（首行为各列名称）")
        row = doc.tables[0].rows[0]
        text = ",".join(cell.text.strip() for cell in row.cells)

    if len(text) > MAX_TEXT_LEN:
        raise ValueError(f"模板文本超长（{len(text)} > {MAX_TEXT_LEN} 字符），请精简后重试")
    return text


def extract_attr_headers(text: str, max_attrs: int = DEFAULT_MAX_ATTRS) -> List[str]:
    """文本 → 表头属性名列表（首个有效行）。

    兼容三种行式：CSV 行（逗号分隔）、md 表格行（`|a|b|c|`，全角竖线
    `｜` 一并宽容）、xlsx 转出的 CSV 文本。跳过空行、`#` 注释行、
    md 分隔行（`|---|:---:|`）；单元格去空白、去空项、去重。
    返回空列表表示没有可解析的表头。
    """
    for line in (text or "").splitlines():
        line = line.strip().replace("｜", "|")
        if not line or line.startswith("#"):
            continue
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            nonempty = [c for c in cells if c]
            # md 分隔行：所有非空单元格仅由 - 与 : 组成
            if nonempty and all(set(c) <= set("-:") for c in nonempty):
                continue
        else:
            try:
                cells = next(csv.reader(io.StringIO(line)))
            except (csv.Error, StopIteration):
                continue
        out: List[str] = []
        for c in cells:
            c = c.strip()
            if c and c not in out:
                out.append(c)
        if out:
            return out[:max_attrs]
    return []


# ===========================================================================
# 模板样式（2026-09-12）：一个栏目 → 多套样式
# ===========================================================================
def style_label(filename: str) -> str:
    """模板文件名 → 样式标签（去扩展名）。空名 → `未命名样式`。

    `label` 是样式在栏目内的**身份键**：同名文件重传 = 同一套（覆盖），
    不同名 = 另一套（新增）。故它必须**只由文件名派生**，不掺时间戳 ——
    否则"改一改再传"会变成新增一份，档案越滚越多。
    """
    stem = Path(str(filename or "")).stem.strip()
    return stem or _UNNAMED


def styles_of(project: Optional[dict]) -> List[dict]:
    """项目的模板样式列表。**零迁移**：只有旧单值字段时合成一条。

    ⚠ 返回的是**拷贝**：调用方拿到后改动，不会回写 `project`（避免"读函数悄悄改数据"）。
    """
    if not isinstance(project, dict):
        return []
    raw = project.get(STYLES_FIELD)
    out: List[dict] = []
    if isinstance(raw, list):
        for s in raw:
            if isinstance(s, dict) and s.get("attrs"):
                d = dict(s)
                d["label"] = d.get("label") or style_label(d.get("filename"))
                out.append(d)
    if out:
        return out
    legacy = project.get(LEGACY_FIELD)
    if isinstance(legacy, dict) and legacy.get("attrs"):
        d = dict(legacy)
        d["label"] = d.get("label") or style_label(d.get("filename"))
        out.append(d)
    return out


def find_style(styles: List[dict], label: str) -> Optional[dict]:
    """按 label 找样式（找不到返回 None）。"""
    for s in styles:
        if s.get("label") == label:
            return s
    return None


def upsert_style(styles: List[dict], style: dict) -> List[dict]:
    """并入一条样式：同 label 覆盖（保位置），否则**追加**。返回新列表。

    "保位置"是刻意的：下拉列表的顺序不该因为"改了一次模板"就跳。
    """
    label = style.get("label") or style_label(style.get("filename"))
    s = dict(style)
    s["label"] = label
    out = [dict(x) for x in styles]
    for i, x in enumerate(out):
        if x.get("label") == label:
            out[i] = s
            return out
    out.append(s)
    return out


def profile_name_for(project_name: str, label: str, *, has_prev: bool) -> str:
    """样式 → 档案名。

    - **该栏目此前没有任何样式** → 直接用栏目名（与 2026-09-06 的旧行为一致，
      故存量栏目重传一次不会被改名、不会多出一个档案）。
    - 已有样式（这是第二套及以后） → `栏目名｜样式标签`，与第一套分开。

    重传**已有 label** 时不走这里 —— 那种情况必须沿用该样式已登记的
    `profile_id`（否则"第一套样式重传"会另建出 `栏目｜A`，出现重名档案）。
    """
    name = str(project_name or "").strip() or "未命名栏目"
    if not has_prev:
        return name
    return f"{name}{LABEL_SEP}{label}"
